import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

doc = docx.Document()

# Page Setup: Margins 0.8 inch
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

def set_cell_shading(cell, color_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_shading(cell, 'F1F5F9')
    set_cell_margins(cell, 120, 120, 200, 200)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="2563EB"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(code_text.strip())
    run.font.name = 'Consolas'
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

# Title
p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_title.paragraph_format.space_after = Pt(2)
r_title = p_title.add_run('TỔNG HỢP CÁC LỆNH DOCKER & DOCKER COMPOSE HỮU DỤNG')
r_title.font.name = 'Arial'
r_title.font.size = Pt(16)
r_title.font.bold = True
r_title.font.color.rgb = RGBColor(30, 58, 138)

p_sub = doc.add_paragraph()
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_sub.paragraph_format.space_after = Pt(14)
r_sub = p_sub.add_run('Hệ Thống Edge-based Open-Set Face Recognition trên NVIDIA Jetson Orin Nano (IP: 10.39.4.131)')
r_sub.font.name = 'Arial'
r_sub.font.size = Pt(10.5)
r_sub.font.italic = True
r_sub.font.color.rgb = RGBColor(100, 116, 139)

# Section 1: Overview
h1 = doc.add_paragraph()
h1.paragraph_format.space_before = Pt(12)
h1.paragraph_format.space_after = Pt(6)
r_h1 = h1.add_run('1. Danh Sách Các Dịch Vụ & Container Trong Hệ Thống')
r_h1.font.name = 'Arial'
r_h1.font.size = Pt(13)
r_h1.font.bold = True
r_h1.font.color.rgb = RGBColor(30, 64, 175)

tbl_services = doc.add_table(rows=5, cols=4)
tbl_services.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ['Tên Container', 'Dịch Vụ', 'Cổng (Port)', 'Vai Trò / Ghi Chú']
for c_idx, text in enumerate(headers):
    cell = tbl_services.cell(0, c_idx)
    set_cell_shading(cell, '1E40AF')
    set_cell_margins(cell, 100, 100, 120, 120)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)

data = [
    ('open_set_fr_backend', 'FastAPI Backend', '8000', 'REST API, WebSocket suy luận, Gallery matrix RAM'),
    ('open_set_fr_frontend', 'Nginx Web UI', '3000', 'Giao diện React dashboard điều khiển & quản trị'),
    ('open_set_fr_postgres', 'PostgreSQL 15', '5432', 'Lưu trữ Database: persons, embeddings, thresholds, users'),
    ('open_set_fr_mediamtx', 'MediaMTX Server', '8554 / 8889', 'Streaming video RTSP / WebRTC đa luồng camera')
]

for r_idx, row_data in enumerate(data, start=1):
    bg_color = 'F8FAFC' if r_idx % 2 == 0 else 'FFFFFF'
    for c_idx, val in enumerate(row_data):
        cell = tbl_services.cell(r_idx, c_idx)
        set_cell_shading(cell, bg_color)
        set_cell_margins(cell, 80, 80, 120, 120)
        p = cell.paragraphs[0]
        run = p.add_run(val)
        run.font.name = 'Arial'
        run.font.size = Pt(9.5)
        if c_idx == 0:
            run.font.name = 'Consolas'
            run.font.bold = True

doc.add_paragraph().paragraph_format.space_after = Pt(8)

# Section 2: Full System Management
h2 = doc.add_paragraph()
h2.paragraph_format.space_before = Pt(10)
h2.paragraph_format.space_after = Pt(4)
r_h2 = h2.add_run('2. Quản Lý Toàn Bộ Hệ Thống (Toàn Bộ Container)')
r_h2.font.name = 'Arial'
r_h2.font.size = Pt(13)
r_h2.font.bold = True
r_h2.font.color.rgb = RGBColor(30, 64, 175)

p_desc2 = doc.add_paragraph('Thực hiện tại thư mục dự án trên Jetson (nơi chứa file docker-compose.yml):')
p_desc2.paragraph_format.space_after = Pt(4)
p_desc2.runs[0].font.size = Pt(10)

add_code_block(doc, """# 1. Khởi động toàn bộ các dịch vụ chạy ngầm (-d)
docker compose up -d

# 2. Xem trạng thái hoạt động (State, Cổng, Thời gian chạy)
docker compose ps

# 3. Khởi động lại toàn bộ hệ thống
docker compose restart

# 4. Tạm dừng toàn bộ dịch vụ (dữ liệu vẫn được giữ nguyên an toàn)
docker compose stop

# 5. Tiếp tục chạy lại các dịch vụ đã tạm dừng
docker compose start

# 6. Tắt hoàn toàn và gỡ bỏ các container
docker compose down""")

# Section 3: Individual Services
h3 = doc.add_paragraph()
h3.paragraph_format.space_before = Pt(10)
h3.paragraph_format.space_after = Pt(4)
r_h3 = h3.add_run('3. Thao Tác Với Từng Dịch Vụ Riêng Biệt')
r_h3.font.name = 'Arial'
r_h3.font.size = Pt(13)
r_h3.font.bold = True
r_h3.font.color.rgb = RGBColor(30, 64, 175)

p_desc3 = doc.add_paragraph('Rất hữu ích khi bạn chỉnh sửa mã nguồn Backend hoặc Frontend mà không muốn làm gián đoạn Database:')
p_desc3.paragraph_format.space_after = Pt(4)
p_desc3.runs[0].font.size = Pt(10)

add_code_block(doc, """# Khởi động lại riêng Backend (áp dụng code mới hoặc nạp lại ma trận nhận diện)
docker compose restart backend

# Khởi động lại riêng Frontend (Web UI)
docker compose restart frontend

# Khởi động lại riêng Database PostgreSQL
docker compose restart postgres

# Dừng riêng một dịch vụ cụ thể
docker compose stop backend

# Bật lại dịch vụ đó
docker compose start backend""")

# Section 4: Logs
h4 = doc.add_paragraph()
h4.paragraph_format.space_before = Pt(10)
h4.paragraph_format.space_after = Pt(4)
r_h4 = h4.add_run('4. Xem Logs & Gỡ Lỗi Realtime (Debugging)')
r_h4.font.name = 'Arial'
r_h4.font.size = Pt(13)
r_h4.font.bold = True
r_h4.font.color.rgb = RGBColor(30, 64, 175)

p_desc4 = doc.add_paragraph('Theo dõi trực tiếp quá trình suy luận, request API hoặc lỗi kết nối (nhấn Ctrl + C để dừng xem log):')
p_desc4.paragraph_format.space_after = Pt(4)
p_desc4.runs[0].font.size = Pt(10)

add_code_block(doc, """# Xem log realtime của TOÀN BỘ hệ thống
docker compose logs -f

# Xem log realtime riêng của Backend (xem luồng API, nhận diện khuôn mặt, EVT)
docker compose logs -f backend

# Xem 100 dòng log gần nhất của Backend rồi mới theo dõi tiếp
docker compose logs --tail 100 -f backend

# Xem log của Database PostgreSQL (kiểm tra truy vấn và kết nối)
docker compose logs -f postgres

# Xem log của Frontend (Nginx web server)
docker compose logs -f frontend""")

# Section 5: Exec Inside Container
h5 = doc.add_paragraph()
h5.paragraph_format.space_before = Pt(10)
h5.paragraph_format.space_after = Pt(4)
r_h5 = h5.add_run('5. Tương Tác Trực Tiếp Bên Trong Container (Docker Exec)')
r_h5.font.name = 'Arial'
r_h5.font.size = Pt(13)
r_h5.font.bold = True
r_h5.font.color.rgb = RGBColor(30, 64, 175)

p_desc5 = doc.add_paragraph('Thao tác trực tiếp với môi trường máy chủ ảo bên trong container:')
p_desc5.paragraph_format.space_after = Pt(4)
p_desc5.runs[0].font.size = Pt(10)

add_code_block(doc, """# 1. Mở cửa sổ dòng lệnh (bash) bên trong container Backend
docker exec -it open_set_fr_backend bash

# 2. Đăng nhập trực tiếp vào PostgreSQL command line (psql)
docker exec -it open_set_fr_postgres psql -U open_set_fr -d open_set_fr

# 3. Chạy nhanh truy vấn đếm số người trong database (không cần vào psql)
docker exec -it open_set_fr_postgres psql -U open_set_fr -d open_set_fr -c "SELECT count(*) FROM persons;"

# 4. Kiểm tra danh sách tài khoản quản trị (users)
docker exec -it open_set_fr_postgres psql -U open_set_fr -d open_set_fr -c "SELECT username, role, is_active FROM users;"

# 5. Chạy nhanh một lệnh Python bên trong môi trường Backend
docker exec -it open_set_fr_backend python3 -c "import torch, cv2; print('Python libraries OK')" """)

# Section 6: Rebuild & Maintenance
h6 = doc.add_paragraph()
h6.paragraph_format.space_before = Pt(10)
h6.paragraph_format.space_after = Pt(4)
r_h6 = h6.add_run('6. Cập Nhật Code Lớn & Dọn Dẹp Bộ Nhớ Jetson')
r_h6.font.name = 'Arial'
r_h6.font.size = Pt(13)
r_h6.font.bold = True
r_h6.font.color.rgb = RGBColor(30, 64, 175)

p_desc6 = doc.add_paragraph('Dùng khi bạn thêm thư viện mới vào requirements.txt hoặc thay đổi Dockerfile:')
p_desc6.paragraph_format.space_after = Pt(4)
p_desc6.runs[0].font.size = Pt(10)

add_code_block(doc, """# 1. Build lại và khởi chạy ngay lập tức
docker compose up -d --build

# 2. Build lại hoàn toàn sạch sẽ (bỏ qua toàn bộ cache cũ)
docker compose build --no-cache
docker compose up -d

# 3. Dọn dẹp các container, image cũ không dùng để giải phóng ổ cứng Jetson
docker system prune -f""")

# Section 7: Project Scripts
h7 = doc.add_paragraph()
h7.paragraph_format.space_before = Pt(10)
h7.paragraph_format.space_after = Pt(4)
r_h7 = h7.add_run('7. Các Script 1-Click Có Sẵn Trong Thư Mục scripts/')
r_h7.font.name = 'Arial'
r_h7.font.size = Pt(13)
r_h7.font.bold = True
r_h7.font.color.rgb = RGBColor(30, 64, 175)

p_desc7 = doc.add_paragraph('Các công cụ tự động hóa đã được thiết lập sẵn trong dự án:')
p_desc7.paragraph_format.space_after = Pt(4)
p_desc7.runs[0].font.size = Pt(10)

add_code_block(doc, """# 1. Khởi động toàn diện hệ thống Docker với host network
bash scripts/run_docker_jetson.sh

# 2. Tạo hoặc reset mật khẩu tài khoản admin (admin / admin)
bash scripts/create_admin.sh

# 3. Quét & di chuyển toàn bộ dữ liệu cũ (ảnh, database, EVT thresholds) sang Docker
bash scripts/migrate_all_legacy_data.sh

# 4. Kiểm tra tổng quan toàn bộ bảng và dữ liệu trong PostgreSQL
python3 scripts/inspect_jetson_db.py""")

# Save to Downloads
downloads_dir = os.path.join(os.environ.get('USERPROFILE', 'C:\\Users\\admin'), 'Downloads')
out_path = os.path.join(downloads_dir, 'Docker_Commands_Huong_Dan_Jetson.docx')
doc.save(out_path)
print(f"SUCCESS: Saved file to {out_path}")
